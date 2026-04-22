"""Literature Scraping Tool - Streamlit App.

This app extracts text from PDF, DOCX, or TXT files and either searches for keywords or extracts all numeric values.
"""

from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional
import re
import tempfile

import PyPDF2
from docx import Document
import streamlit as st


@dataclass
class PaperInfo:
    path: Path
    text: str
    page_count: int
    file_size_bytes: int
    title: Optional[str]
    author: Optional[str]
    producer: Optional[str]
    creation_date: Optional[str]
    modification_date: Optional[str]
    character_count: int
    word_count: int
    metadata: Dict[str, str]


def format_bytes(size: int) -> str:
    """Format bytes into a human-readable string."""
    for unit in ["B", "KB", "MB", "GB"]:
        if size < 1024.0:
            return f"{size:.1f} {unit}"
        size /= 1024.0
    return f"{size:.1f} TB"


def extract_text_from_file(file_input) -> PaperInfo:
    """Extract plain text and metadata from a PDF, DOCX, or TXT file.
    
    Args:
        file_input: Either a Path object or a file-like object from st.file_uploader
    """
    # Determine file type and handle both Path and file-like objects
    if isinstance(file_input, Path):
        suffix = file_input.suffix.lower()
        file_name = file_input.name
        file_size_bytes = file_input.stat().st_size
        file_obj = None
        is_path = True
    else:
        # File-like object from Streamlit
        suffix = Path(file_input.name).suffix.lower()
        file_name = file_input.name
        file_size_bytes = len(file_input.getvalue()) if hasattr(file_input, 'getvalue') else 0
        file_obj = file_input
        is_path = False
    
    text = ""
    metadata = {}
    page_count = 0

    if suffix == ".pdf":
        if is_path:
            with file_input.open("rb") as f:
                reader = PyPDF2.PdfReader(f)
                pages = reader.pages
                text_chunks: List[str] = [page.extract_text() or "" for page in pages]
                text = "\n".join(text_chunks)
                raw_metadata = reader.metadata or {}
                metadata = {k[1:]: str(v) for k, v in raw_metadata.items() if k.startswith("/") and v is not None}
                page_count = len(pages)
        else:
            reader = PyPDF2.PdfReader(file_obj)
            pages = reader.pages
            text_chunks: List[str] = [page.extract_text() or "" for page in pages]
            text = "\n".join(text_chunks)
            raw_metadata = reader.metadata or {}
            metadata = {k[1:]: str(v) for k, v in raw_metadata.items() if k.startswith("/") and v is not None}
            page_count = len(pages)
    elif suffix == ".docx":
        if is_path:
            doc = Document(file_input)
        else:
            doc = Document(file_obj)
        text_chunks = [para.text for para in doc.paragraphs if para.text.strip()]
        text = "\n".join(text_chunks)
        # Basic metadata from docx
        core_props = doc.core_properties
        metadata = {
            "Title": core_props.title,
            "Author": core_props.author,
            "Created": str(core_props.created) if core_props.created else None,
            "Modified": str(core_props.modified) if core_props.modified else None,
        }
        metadata = {k: v for k, v in metadata.items() if v}
        page_count = 1  # DOCX doesn't have pages easily
    elif suffix == ".txt":
        if is_path:
            with file_input.open("r", encoding="utf-8") as f:
                text = f.read()
        else:
            text = file_obj.getvalue().decode("utf-8")
        page_count = 1  # TXT doesn't have pages
    else:
        raise ValueError(f"Unsupported file type: {suffix}")

    if not text:
        text = ""

    return PaperInfo(
        path=Path(file_name),
        text=text,
        page_count=page_count,
        file_size_bytes=file_size_bytes,
        title=metadata.get("Title"),
        author=metadata.get("Author"),
        producer=metadata.get("Producer"),
        creation_date=metadata.get("Created") or metadata.get("CreationDate"),
        modification_date=metadata.get("Modified") or metadata.get("ModDate"),
        character_count=len(text),
        word_count=len(text.split()),
        metadata=metadata,
    )


def setup_streamlit_ui():
    """Set up the Streamlit user interface and return user inputs."""
    st.set_page_config(page_title="Literature Scraper", layout="wide")
    st.title("📄 Literature Scraper")
    
    st.write("Upload up to 10 PDF, DOCX, or TXT files and extract information.")
    
    # File upload
    uploaded_files = st.file_uploader(
        "Upload files",
        type=["pdf", "docx", "txt"],
        accept_multiple_files=True,
        key="file_uploader"
    )
    
    if len(uploaded_files) > 10:
        st.error("Please upload no more than 10 files.")
        return None, None, None
    
    # Mode selection
    mode = st.radio(
        "Choose mode:",
        ["Search Keywords", "Extract Numbers"],
        key="mode_radio"
    )
    
    # Keywords input (conditional)
    keywords = None
    if mode == "Search Keywords":
        keywords_text = st.text_input(
            "Enter keywords to search for (separated by commas):",
            key="keywords_input"
        )
        if keywords_text:
            keywords = [k.strip() for k in keywords_text.split(",")]
    
    # Process button
    if st.button("Process Files", type="primary"):
        if not uploaded_files:
            st.warning("Please upload at least one file.")
            return None, None, None
        if mode == "Search Keywords" and not keywords:
            st.warning("Please enter keywords.")
            return None, None, None
        return uploaded_files, mode, keywords
    
    return None, None, None


def search_keywords(text: str, keywords: List[str]) -> Dict[str, int]:
    """Search for keywords in text and return counts (case-insensitive)."""
    text_lower = text.lower()
    counts = {}
    for keyword in keywords:
        keyword_lower = keyword.lower()
        count = text_lower.count(keyword_lower)
        counts[keyword] = count
    return counts


def extract_numbers(text: str) -> List[str]:
    """Extract all numeric values from text using regex (integers, decimals, fractions, complex)."""
    # Regex patterns for different number types
    patterns = [
        r'\b\d+\.\d+\b',  # decimals like 3.14
        r'\b\d+/\d+\b',   # fractions like 1/2
        r'\b\d+\+\d+i\b', # complex like 2+3i
        r'\b\d+i\b',      # imaginary like 5i
        r'\b\d+\b',       # integers like 42
    ]
    numbers = []
    for pattern in patterns:
        matches = re.findall(pattern, text)
        numbers.extend(matches)
    # Remove duplicates while preserving order
    seen = set()
    unique_numbers = []
    for num in numbers:
        if num not in seen:
            unique_numbers.append(num)
            seen.add(num)
    return unique_numbers


def format_paper_summary(info: PaperInfo, keyword_counts: Optional[Dict[str, int]] = None, numbers: Optional[List[str]] = None) -> str:
    """Return a formatted summary of paper attributes and keyword matches or extracted numbers."""
    metadata_lines = []
    field_mapping = {
        "Title": "title",
        "Author": "author",
        "Producer": "producer",
        "CreationDate": "creation_date",
        "ModDate": "modification_date",
    }

    for display_name, attr_name in field_mapping.items():
        value = getattr(info, attr_name, None)
        if value:
            metadata_lines.append(f"    {display_name}: {value}")

    custom_metadata = {
        k: v for k, v in info.metadata.items() if k not in {"Title", "Author", "Producer", "CreationDate", "ModDate"}
    }
    if custom_metadata:
        metadata_lines.append("    Other metadata:")
        for key, value in custom_metadata.items():
            metadata_lines.append(f"      {key}: {value}")

    preview = info.text.strip().replace("\n", " ")[:200]
    if len(info.text.strip()) > 200:
        preview += "..."

    result_lines = []
    if keyword_counts is not None:
        total_matches = sum(keyword_counts.values())
        if total_matches > 0:
            result_lines.append("  Keyword matches:")
            for keyword, count in keyword_counts.items():
                if count > 0:
                    result_lines.append(f"    '{keyword}': {count} times")
        else:
            result_lines.append("  No keyword matches found.")
    elif numbers is not None:
        if numbers:
            result_lines.append(f"  Extracted numbers ({len(numbers)} found):")
            result_lines.append(f"    {', '.join(numbers[:20])}")  # Limit to first 20 for brevity
            if len(numbers) > 20:
                result_lines.append(f"    ... and {len(numbers) - 20} more")
        else:
            result_lines.append("  No numbers found.")

    return (
        f"Paper: {info.path.name}\n"
        f"  Path: {info.path.resolve()}\n"
        f"  File size: {format_bytes(info.file_size_bytes)}\n"
        f"  Pages: {info.page_count}\n"
        f"  Characters extracted: {info.character_count}\n"
        f"  Words extracted: {info.word_count}\n"
        + ("\n".join(metadata_lines) + "\n" if metadata_lines else "")
        + "\n".join(result_lines) + "\n"
        + f"  Preview: {preview}\n"
    )


def main() -> None:
    uploaded_files, mode, keywords = setup_streamlit_ui()
    
    if uploaded_files is None:
        return

    # Process files
    results = []
    for uploaded_file in uploaded_files:
        try:
            info = extract_text_from_file(uploaded_file)
            
            if mode == "Search Keywords":
                keyword_counts = search_keywords(info.text, keywords)
                total_matches = sum(keyword_counts.values())
                results.append((info, keyword_counts, None))
            else:  # Extract Numbers
                numbers = extract_numbers(info.text)
                results.append((info, None, numbers))
        except Exception as e:
            st.error(f"Error processing {uploaded_file.name}: {str(e)}")
            continue

    # Display results
    if results:
        st.divider()
        st.subheader("Results")
        
        for info, keyword_counts, numbers in results:
            with st.expander(f"📋 {info.path.name}", expanded=True):
                # Display summary in columns
                col1, col2, col3 = st.columns(3)
                with col1:
                    st.metric("Pages", info.page_count)
                with col2:
                    st.metric("Words", info.word_count)
                with col3:
                    st.metric("Size", format_bytes(info.file_size_bytes))
                
                # Metadata
                if info.title or info.author or info.producer:
                    st.write("**Metadata:**")
                    metadata_dict = {}
                    if info.title:
                        metadata_dict["Title"] = info.title
                    if info.author:
                        metadata_dict["Author"] = info.author
                    if info.producer:
                        metadata_dict["Producer"] = info.producer
                    if info.creation_date:
                        metadata_dict["Created"] = info.creation_date
                    if info.modification_date:
                        metadata_dict["Modified"] = info.modification_date
                    st.json(metadata_dict)
                
                # Results
                if keyword_counts is not None:
                    total_matches = sum(keyword_counts.values())
                    st.write(f"**Keyword Matches** ({total_matches} total):")
                    keyword_dict = {k: v for k, v in keyword_counts.items() if v > 0}
                    if keyword_dict:
                        st.json(keyword_dict)
                    else:
                        st.write("No matches found")
                
                elif numbers is not None:
                    st.write(f"**Extracted Numbers** ({len(numbers)} found):")
                    numbers_display = ", ".join(numbers[:50])
                    if len(numbers) > 50:
                        numbers_display += f"\n... and {len(numbers) - 50} more"
                    st.text(numbers_display)
                
                # Preview
                preview = info.text.strip().replace("\n", " ")[:300]
                if len(info.text.strip()) > 300:
                    preview += "..."
                st.write("**Preview:**")
                st.text(preview)
        
        # Summary
        st.divider()
        st.subheader("Summary")
        if mode == "Search Keywords":
            matching_count = sum(1 for _, keyword_counts, _ in results if keyword_counts and sum(keyword_counts.values()) > 0)
            st.write(f"✅ {matching_count} out of {len(results)} files contain the keywords")
        else:
            st.write(f"✅ Processed {len(results)} files successfully")


if __name__ == "__main__":
    main()
    


if __name__ == "__main__":
    main()
    